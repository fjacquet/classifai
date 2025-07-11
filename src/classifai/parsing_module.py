"""
Parsing module for ClassifAI.

This module encapsulates the logic for extracting content from various file types
using a hierarchical strategy.
"""

import email
import os
import subprocess
import tarfile
import tempfile
import zipfile

import extract_msg
import fitz  # PyMuPDF
import openpyxl
import pytesseract
from bs4 import BeautifulSoup
from docx import Document
from loguru import logger
from PIL import Image, UnidentifiedImageError
from PIL.ExifTags import TAGS
from striprtf.striprtf import rtf_to_text

from classifai.config import GENERIC_TEXT_EXTENSIONS
from classifai.geocoding_module import get_location_from_gps


def _get_exif_data(image: Image) -> dict:
    """Extracts and decodes EXIF data from an image."""
    exif_data = {}
    if hasattr(image, "_getexif"):
        exif_info = image._getexif()
        if exif_info:
            for tag, value in exif_info.items():
                decoded = TAGS.get(tag, tag)
                exif_data[decoded] = value
    return exif_data


def _convert_gps_to_decimal(gps_coords, gps_ref):
    """Converts GPS coordinates from DMS to decimal degrees."""
    decimal_degrees = gps_coords[0] + (gps_coords[1] / 60) + (gps_coords[2] / 3600)
    if gps_ref in ["S", "W"]:
        decimal_degrees = -decimal_degrees
    return decimal_degrees


def parse_image(file_path: str) -> tuple[str, dict]:
    """
    Extracts text and metadata from an image file.

    Returns:
        A tuple containing the text content and a dictionary of metadata.
    """
    metadata = {}
    try:
        with Image.open(file_path) as img:
            # Extract text with OCR
            text = pytesseract.image_to_string(img)
            if not text.strip():
                logger.info(f"OCR returned no text for image: {file_path}")

            # Extract EXIF data
            exif_data = _get_exif_data(img)
            if "DateTimeOriginal" in exif_data:
                metadata["date"] = exif_data["DateTimeOriginal"]
            if "GPSInfo" in exif_data:
                gps_info = exif_data["GPSInfo"]
                lat = _convert_gps_to_decimal(gps_info[2], gps_info[1])
                lon = _convert_gps_to_decimal(gps_info[4], gps_info[3])
                metadata["location"] = get_location_from_gps(lat, lon)

            return text, metadata

    except (UnidentifiedImageError, ValueError):
        logger.warning(f"Cannot identify image file: {file_path}")
        return "", {}
    except pytesseract.TesseractNotFoundError:
        logger.error("Tesseract is not installed or not in your PATH.")
        raise
    except Exception as e:
        logger.error(f"Error performing OCR on image file {file_path}: {e}")
        return "", {}


def parse_pdf(file_path: str) -> tuple[str, dict]:
    """Extracts text content from a PDF file."""
    try:
        with fitz.open(file_path) as doc:
            text = "".join(page.get_text() for page in doc)
        if not text.strip():
            logger.info(f"PyMuPDF found no text in {file_path}. Trying pdftotext.")
            try:
                result = subprocess.run(
                    ["pdftotext", file_path, "-"],
                    capture_output=True,
                    text=True,
                    check=True,
                )
                return result.stdout, {}
            except (subprocess.CalledProcessError, FileNotFoundError) as e:
                logger.warning(
                    f"pdftotext fallback failed for {file_path}: {e}. Ensure 'poppler-utils' is installed."
                )
                return "", {}
        return text, {}
    except Exception as e:
        logger.error(f"Error parsing PDF file {file_path}: {e}")
        return "", {}


def parse_docx(file_path: str) -> tuple[str, dict]:
    """Extracts text content from a DOCX file."""
    try:
        doc = Document(file_path)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs]), {}
    except Exception as e:
        logger.error(f"Error parsing DOCX file {file_path}: {e}")
        return "", {}


def parse_xlsx(file_path: str) -> tuple[str, dict]:
    """Extracts text content from an XLSX file."""
    try:
        workbook = openpyxl.load_workbook(file_path)
        text = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        text.append(str(cell.value))
        return "\n".join(text), {}
    except Exception as e:
        logger.error(f"Error parsing XLSX file {file_path}: {e}")
        return "", {}


def parse_html(file_path: str) -> tuple[str, dict]:
    """Extracts text from an HTML file."""
    try:
        with open(file_path, encoding="utf-8") as f:
            soup = BeautifulSoup(f, "html.parser")
            return soup.get_text(), {}
    except Exception as e:
        logger.error(f"Error parsing HTML file {file_path}: {e}")
        return "", {}


def parse_rtf(file_path: str) -> tuple[str, dict]:
    """Extracts text from an RTF file."""
    try:
        with open(file_path) as f:
            return rtf_to_text(f.read()), {}
    except Exception as e:
        logger.error(f"Error parsing RTF file {file_path}: {e}")
        return "", {}


def parse_eml(file_path: str) -> tuple[str, dict]:
    """Extracts text from an EML file."""
    try:
        with open(file_path) as f:
            msg = email.message_from_file(f)
            body = ""
            if msg.is_multipart():
                for part in msg.walk():
                    if part.get_content_type() == "text/plain":
                        body = part.get_payload(decode=True).decode()
                        break
            else:
                body = msg.get_payload(decode=True).decode()
            return body, {}
    except Exception as e:
        logger.error(f"Error parsing EML file {file_path}: {e}")
        return "", {}


def parse_msg(file_path: str) -> tuple[str, dict]:
    """Extracts text from an MSG file."""
    try:
        with extract_msg.openMsg(file_path) as msg:
            return msg.body, {}
    except Exception as e:
        logger.error(f"Error parsing MSG file {file_path}: {e}")
        return "", {}


def parse_generic_text(file_path: str) -> tuple[str, dict]:
    """Extracts content from a generic text file, trying various encodings."""
    encodings = ["utf-8", "latin-1", "iso-8859-1"]
    for encoding in encodings:
        try:
            with open(file_path, encoding=encoding) as f:
                return f.read(), {}
        except UnicodeDecodeError:
            continue
    logger.warning(f"Could not decode file {file_path} with any of the default encodings.")
    return "", {}


def parse_with_pandoc(file_path: str) -> tuple[str, dict]:
    """Uses Pandoc as a fallback to convert a document to plain text."""
    try:
        result = subprocess.run(
            ["pandoc", file_path, "-t", "plain"],
            capture_output=True,
            text=True,
            check=True,
        )
        return result.stdout, {}
    except FileNotFoundError:
        logger.error(
            "Pandoc is not installed or not in your PATH. Please install it for advanced document parsing."
        )
        # Re-raise to avoid repeated attempts
        raise
    except subprocess.CalledProcessError as e:
        logger.warning(f"Pandoc failed to parse {file_path}: {e.stderr}")
        return "", {}


def parse_archive(file_path: str) -> tuple[str, dict]:
    """
    Extracts content from files within a ZIP or TAR archive by extracting them
    to a temporary directory and parsing them individually.
    """
    text_content = []
    file_path_lower = file_path.lower()

    with tempfile.TemporaryDirectory() as temp_dir:
        try:
            if file_path_lower.endswith(".zip"):
                with zipfile.ZipFile(file_path, "r") as archive:
                    archive.extractall(temp_dir)
                    for root, _, files in os.walk(temp_dir):
                        for name in files:
                            file_ext = "." + name.split(".")[-1]
                            parser = get_parser(file_ext)
                            if parser and parser != parse_archive:
                                temp_file_path = os.path.join(root, name)
                                content, _ = parser(temp_file_path)
                                if content:
                                    text_content.append(f"--- Content from {name} ---\n{content}")

            elif file_path_lower.endswith((".tar", ".gz", ".bz2", ".xz")):
                with tarfile.open(file_path, "r:*") as archive:
                    archive.extractall(temp_dir)
                    for root, _, files in os.walk(temp_dir):
                        for name in files:
                            file_ext = "." + name.split(".")[-1]
                            parser = get_parser(file_ext)
                            if parser and parser != parse_archive:
                                temp_file_path = os.path.join(root, name)
                                content, _ = parser(temp_file_path)
                                if content:
                                    text_content.append(f"--- Content from {name} ---\n{content}")

            return "\n\n".join(text_content), {}

        except (zipfile.BadZipFile, tarfile.ReadError) as e:
            logger.error(f"Could not read archive {file_path}: {e}")
            return "", {}
        except Exception as e:
            logger.error(f"Error parsing archive {file_path}: {e}")
            return "", {}


def get_parser(file_extension: str):
    """
    Returns the appropriate parser function for a given file extension using
    a hierarchical strategy.
    """
    specific_parsers = {
        # Text & Documents
        ".pdf": parse_pdf,
        ".docx": parse_docx,
        ".xlsx": parse_xlsx,
        ".html": parse_html,
        ".htm": parse_html,
        ".rtf": parse_rtf,
        ".eml": parse_eml,
        ".msg": parse_msg,
        # Images
        ".png": parse_image,
        ".jpg": parse_image,
        ".jpeg": parse_image,
        ".tiff": parse_image,
        ".bmp": parse_image,
        # Archives
        ".zip": parse_archive,
        ".tar": parse_archive,
        ".gz": parse_archive,
        ".bz2": parse_archive,
        ".xz": parse_archive,
    }
    if file_extension in specific_parsers:
        return specific_parsers[file_extension]

    if file_extension in GENERIC_TEXT_EXTENSIONS:
        return parse_generic_text

    if not file_extension:
        return parse_generic_text

    # As a last resort for document-like files, try Pandoc
    pandoc_supported = [".odt", ".epub", ".md"]  # Add more as needed
    if file_extension in pandoc_supported:
        return parse_with_pandoc

    return None
