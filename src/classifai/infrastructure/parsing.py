"""Impure parsing functions for ClassifAI."""

import email
import functools
import os
import stat
import subprocess
import tarfile
import tempfile
import zipfile

import extract_msg

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

import openpyxl
import pytesseract
from bs4 import BeautifulSoup
from docx import Document
from loguru import logger
from PIL import Image, UnidentifiedImageError
from PIL.ExifTags import TAGS
from striprtf.striprtf import rtf_to_text

from classifai.infrastructure.geocoding import get_location_from_gps


# --- Decorator for Exception Handling ---
def parsing_handler(func):
    """
    A decorator to handle common exceptions for parser functions.
    It logs errors and ensures the function returns a default value.
    """

    @functools.wraps(func)
    def wrapper(file_path: str, *args, **kwargs) -> tuple[str, dict]:
        try:
            return func(file_path, *args, **kwargs)
        except (UnidentifiedImageError, ValueError) as e:
            logger.warning(f"Cannot identify or process file '{file_path}': {e}")
        except pytesseract.TesseractNotFoundError:
            logger.error("Tesseract is not installed or not in your PATH.")
            raise  # Re-raise as this is a critical setup issue
        except FileNotFoundError:
            logger.error(f"File not found: {file_path}")
        except (zipfile.BadZipFile, tarfile.ReadError) as e:
            logger.error(f"Could not read archive {file_path}: {e}")
        except Exception as e:
            logger.error(f"An unexpected error occurred in '{func.__name__}' for file '{file_path}': {e}")
        return "", {}

    return wrapper


# --- Specific Parsers ---
def _get_exif_data(image: Image) -> dict:
    """Extracts and decodes EXIF data from an image."""
    exif_data = {}
    if hasattr(image, "_getexif"):
        exif_info = image._getexif()
        if exif_info:
            for tag, value in exif_info.items():
                decoded = TAGS.get(tag, tag)
                exif_data[decoded] = value
    return exif_data


def _convert_gps_to_decimal(gps_coords, gps_ref):
    """Converts GPS coordinates from DMS to decimal degrees."""
    decimal_degrees = gps_coords[0] + (gps_coords[1] / 60) + (gps_coords[2] / 3600)
    if gps_ref in ["S", "W"]:
        decimal_degrees = -decimal_degrees
    return decimal_degrees


@parsing_handler
def parse_image(file_path: str) -> tuple[str, dict]:
    """Extracts text and metadata from an image file."""
    metadata = {}
    with Image.open(file_path) as img:
        text = pytesseract.image_to_string(img)
        if not text.strip():
            logger.info(f"OCR returned no text for image: {file_path}")

        exif_data = _get_exif_data(img)
        if "DateTimeOriginal" in exif_data:
            metadata["date"] = exif_data["DateTimeOriginal"]
        if "GPSInfo" in exif_data:
            gps_info = exif_data["GPSInfo"]
            lat = _convert_gps_to_decimal(gps_info[2], gps_info[1])
            lon = _convert_gps_to_decimal(gps_info[4], gps_info[3])
            metadata["location"] = get_location_from_gps(lat, lon)

        return text, metadata


@parsing_handler
def parse_pdf(file_path: str) -> tuple[str, dict]:
    """Extracts text content from a PDF file."""
    if not fitz:
        logger.warning("PyMuPDF (fitz) is not installed. Falling back to pdftotext.")
        result = subprocess.run(["pdftotext", file_path, "-"], capture_output=True, text=True, check=True)
        return result.stdout, {}

    with fitz.open(file_path) as doc:
        text = "".join(page.get_text() for page in doc)
    if not text.strip():
        logger.info(f"PyMuPDF found no text in {file_path}. Trying pdftotext.")
        result = subprocess.run(["pdftotext", file_path, "-"], capture_output=True, text=True, check=True)
        return result.stdout, {}
    return text, {}


@parsing_handler
def parse_docx(file_path: str) -> tuple[str, dict]:
    """Extracts text content from a DOCX file."""
    doc = Document(file_path)
    return "\n".join([paragraph.text for paragraph in doc.paragraphs]), {}


@parsing_handler
def parse_xlsx(file_path: str) -> tuple[str, dict]:
    """Extracts text content from an XLSX file."""
    workbook = openpyxl.load_workbook(file_path)
    text = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value:
                    text.append(str(cell.value))
    return "\n".join(text), {}


@parsing_handler
def parse_html(file_path: str) -> tuple[str, dict]:
    """Extracts text from an HTML file."""
    with open(file_path, encoding="utf-8") as f:
        soup = BeautifulSoup(f, "html.parser")
        return soup.get_text(), {}


@parsing_handler
def parse_rtf(file_path: str) -> tuple[str, dict]:
    """Extracts text from an RTF file."""
    with open(file_path) as f:
        return rtf_to_text(f.read()), {}


@parsing_handler
def parse_eml(file_path: str) -> tuple[str, dict]:
    """Extracts text from an EML file."""
    with open(file_path) as f:
        msg = email.message_from_file(f)
        body = ""
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    payload = part.get_payload(decode=True)
                    if isinstance(payload, bytes):
                        body = payload.decode()
                    break
        else:
            payload = msg.get_payload(decode=True)
            if isinstance(payload, bytes):
                body = payload.decode()
        return body, {}


@parsing_handler
def parse_msg(file_path: str) -> tuple[str, dict]:
    """Extracts text and metadata from an MSG file."""
    with extract_msg.openMsg(file_path) as msg:
        # Extract basic content
        body = msg.body or ""

        # Extract metadata
        metadata = {
            "subject": msg.subject or "",
            "sender": msg.sender or "",
            "date": msg.date or "",
            "to": msg.to or "",
            "cc": msg.cc or "",
        }

        # Add attachment information if available
        if msg.attachments:
            attachment_names = [att.longFilename or att.shortFilename for att in msg.attachments]
            metadata["attachments"] = ", ".join(attachment_names)

        # Combine header information with body for better context
        header = f"Subject: {metadata['subject']}\nFrom: {metadata['sender']}\nTo: {metadata['to']}\nDate: {metadata['date']}\n\n"
        full_content = header + body

        logger.debug(f"Extracted metadata from MSG file: {metadata}")
        return full_content, metadata


@parsing_handler
def parse_generic_text(file_path: str) -> tuple[str, dict]:
    """Extracts content from a generic text file, trying various encodings."""
    encodings = ["utf-8", "latin-1", "iso-8859-1"]
    for encoding in encodings:
        try:
            with open(file_path, encoding=encoding) as f:
                return f.read(), {}
        except UnicodeDecodeError:
            continue
    logger.warning(f"Could not decode file {file_path} with any of the default encodings.")
    return "", {}


@parsing_handler
def parse_with_pandoc(file_path: str) -> tuple[str, dict]:
    """Uses Pandoc as a fallback to convert a document to plain text."""
    try:
        result = subprocess.run(
            ["pandoc", file_path, "-t", "plain"],
            capture_output=True,
            text=True,
            check=True,
        )
        return result.stdout, {}
    except FileNotFoundError:
        logger.error(
            "Pandoc is not installed or not in your PATH. Please install it for advanced document parsing.",
        )
        raise
    except subprocess.CalledProcessError as e:
        logger.warning(f"Pandoc failed to parse {file_path}: {e.stderr}")
        return "", {}


@parsing_handler
def parse_archive(file_path: str) -> tuple[str, dict]:
    """
    Extracts the names of files within a ZIP or TAR archive.
    """
    text_content = []
    file_path_lower = file_path.lower()

    def _is_within_dir(base: str | os.PathLike, target: str | os.PathLike) -> bool:
        base_path = os.path.realpath(base)
        target_path = os.path.realpath(target)
        return os.path.commonpath([base_path]) == os.path.commonpath([base_path, target_path])

    def _safe_extract_tar(tf: tarfile.TarFile, dest: str) -> None:
        for member in tf.getmembers():
            # Normalize and validate path
            name = member.name
            norm_name = os.path.normpath(name)
            if os.path.isabs(norm_name) or norm_name.startswith(".." + os.sep) or ".." + os.sep in norm_name:
                logger.warning(f"Blocked unsafe tar member path: {name}")
                continue

            dest_path = os.path.join(dest, norm_name)
            if not _is_within_dir(dest, dest_path):
                logger.warning(f"Blocked traversal outside dest for tar member: {name}")
                continue

            # Directories
            if member.isdir():
                os.makedirs(dest_path, exist_ok=True)
                continue

            # Block symlinks and hardlinks
            if member.issym() or member.islnk():
                logger.warning(f"Skipping link in tar archive: {name}")
                continue

            # Skip special files (devices, fifos, etc.)
            if not member.isreg():
                logger.warning(f"Skipping non-regular tar member: {name}")
                continue

            # Ensure directory exists
            os.makedirs(os.path.dirname(dest_path), exist_ok=True)

            # Extract file content safely without following symlinks
            src = tf.extractfile(member)
            if src is None:
                logger.warning(f"Could not read tar member (None): {name}")
                continue
            with src as s, open(dest_path, "wb") as f:
                while True:
                    chunk = s.read(1024 * 1024)
                    if not chunk:
                        break
                    f.write(chunk)

    def _safe_extract_zip(zf: zipfile.ZipFile, dest: str) -> None:
        for member in zf.infolist():
            name = member.filename
            norm_name = os.path.normpath(name)
            if os.path.isabs(norm_name) or norm_name.startswith(".." + os.sep) or ".." + os.sep in norm_name:
                logger.warning(f"Blocked unsafe zip member path: {name}")
                continue

            dest_path = os.path.join(dest, norm_name)
            if not _is_within_dir(dest, dest_path):
                logger.warning(f"Blocked traversal outside dest for zip member: {name}")
                continue

            # Detect symlink in zip (POSIX) via external attributes
            is_symlink = False
            try:
                external_attr = member.external_attr >> 16
                is_symlink = stat.S_ISLNK(external_attr)
            except Exception:
                is_symlink = False

            if is_symlink:
                logger.warning(f"Skipping symlink in zip archive: {name}")
                continue

            if name.endswith("/") or member.is_dir():
                os.makedirs(dest_path, exist_ok=True)
                continue

            os.makedirs(os.path.dirname(dest_path), exist_ok=True)
            with zf.open(member, "r") as s, open(dest_path, "wb") as f:
                while True:
                    chunk = s.read(1024 * 1024)
                    if not chunk:
                        break
                    f.write(chunk)

    with tempfile.TemporaryDirectory() as temp_dir:
        if file_path_lower.endswith(".zip"):
            with zipfile.ZipFile(file_path, "r") as zip_archive:
                _safe_extract_zip(zip_archive, temp_dir)
        elif file_path_lower.endswith((".tar", ".gz", ".bz2", ".xz")):
            with tarfile.open(file_path, "r:*") as tar_archive:
                _safe_extract_tar(tar_archive, temp_dir)
        else:
            return "", {}

        for _root, _, files in os.walk(temp_dir):
            for name in files:
                text_content.append(f"--- File: {name} ---")

        return "\n".join(text_content), {}
